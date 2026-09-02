"""Flow-document export: turn a translated IR into an editable .docx, and from
there into a Hancom .hwpx (via the bundled docx→hwpx converter).

Unlike the PDF renderer — which edits the original page in place to preserve the
fixed layout — this produces a *flowing* word-processor document: title, headings,
paragraphs and list items in reading order at a uniform body size, ruled tables
rebuilt as real tables, a 2-column body where the paper is two-column, and figures
(vector or raster) rasterized from the page and embedded as pictures. That is the
natural shape for an editable Korean (.hwpx) or Word (.docx) file.
"""
from __future__ import annotations

import contextlib
import importlib
import re
import sys
import tempfile
import threading
from pathlib import Path

import pymupdf as fitz
from docx import Document as Docx
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.shared import Emu, Pt

from .ir import Document, TranslationResult
from .renderer import _strip_math_delims

_PAGE_NUM_RE = re.compile(r"^[\-–—\s]*\d{1,4}[\-–—\s]*$")   # "- 6 -", "12"
_SUBHEAD_RE = re.compile(r"^\s*(?:\d+\.\d+|[가-힣]\.|[a-z]\)|\([0-9가-힣a-z]\))")
# start of the body — first numbered section heading ("1. Introduction", "1 서론")
_BODY_START_RE = re.compile(r"^\s*1\s*[.．]?\s+\S")


def _frontmatter_cutoff(page) -> float | None:
    """y where the body begins on the first page (first "1. …" heading). Content
    above it — title, authors, highlights, abstract — should read as one column,
    not be split into the body's two columns."""
    ys = [b.bbox[1] for b in page.blocks
          if _BODY_START_RE.match((b.text() or "").strip()) and b.bbox[1] > 120]
    return min(ys) if ys else None


# --------------------------------------------------------------------------- #
# Table reconstruction
#
# The IR has no table model — a table's cells arrive as individual text blocks.
# To rebuild an editable table we take the column boundaries + top edge from
# PyMuPDF's table finder (reliable for the ruled header) and derive the body
# rows from the page's full-width horizontal rules, then drop each translated
# block into the (row, col) its position selects. Best-effort: merged/multi-line
# source cells may land together, but the result is a real table, not plain text.
# --------------------------------------------------------------------------- #
def _horizontal_rules(page, x0: float, x1: float) -> list[float]:
    span = x1 - x0
    ys: set[float] = set()
    for dr in page.get_drawings():
        for it in dr.get("items", []):
            if it[0] == "l":
                p1, p2 = it[1], it[2]
                if abs(p1.y - p2.y) < 1.2:
                    a, b = min(p1.x, p2.x), max(p1.x, p2.x)
                    if b - a >= 0.55 * span and a <= x0 + 20 and b >= x1 - 20:
                        ys.add(round(p1.y, 1))
            elif it[0] == "re":
                r = it[1]
                if r.x1 - r.x0 >= 0.55 * span and r.x0 <= x0 + 20 and r.x1 >= x1 - 20:
                    ys.add(round(r.y0, 1))
                    ys.add(round(r.y1, 1))
    return sorted(ys)


def _page_tables(page) -> list[dict]:
    """Return table specs {cols: [x…], rows: [y…], bbox} for a PDF page."""
    specs: list[dict] = []
    try:
        tabs = list(page.find_tables().tables)
    except Exception:
        return specs
    for t in tabs:
        try:
            cells = t.cells
            cols = sorted({round(c[0], 1) for c in cells} | {round(c[2], 1) for c in cells})
            if len(cols) < 3:                       # need ≥ 2 columns
                continue
            x0, x1 = cols[0], cols[-1]
            top = float(t.bbox[1])
            rules = [y for y in _horizontal_rules(page, x0, x1) if y >= top - 2]
            if len(rules) < 2:                      # no real internal row separators
                continue                            # → likely a bordered paragraph
            rows = [round(top, 1)]
            for y in rules:
                gap = y - rows[-1]
                if gap < 2:                         # duplicate rule
                    continue
                if gap > 95:                        # a real section break; a tall
                    break                           # multi-line cell stays smaller
                rows.append(y)
            if len(rows) < 3:                       # need ≥ 2 body rows
                continue
            specs.append({"x0": x0, "x1": x1, "y0": rows[0], "y1": rows[-1],
                          "cols": cols, "rows": rows})
        except Exception:
            continue
    return specs


def _band(rows: list[float], yc: float) -> int | None:
    for i in range(len(rows) - 1):
        if rows[i] - 2 <= yc < rows[i + 1] + 2:
            return i
    return None


# --------------------------------------------------------------------------- #
# Borderless-table reconstruction (no ruling lines)
#
# A borderless table's cells are stored as individual PDF "lines" (each cell is a
# short line at its own x, cells sharing a y form a row). We cluster those cell
# positions — y into rows, x into columns — to rebuild the grid for the flowing
# DOCX/HWPX, which otherwise gets a jumbled single line per row.
# --------------------------------------------------------------------------- #
_TABLE_CAP_RE_DOCX = re.compile(r"^\s*(?:table|tab\.?|표)\s*\d", re.I)


def _cluster_centers(vals: list[float], gap: float) -> list[float]:
    vals = sorted(vals)
    if not vals:
        return []
    groups = [[vals[0]]]
    for v in vals[1:]:
        if v - groups[-1][-1] <= gap:
            groups[-1].append(v)
        else:
            groups.append([v])
    return [sum(g) / len(g) for g in groups]


def _nearest(centers: list[float], v: float) -> int:
    return min(range(len(centers)), key=lambda i: abs(centers[i] - v))


def _borderless_grids(page) -> list[dict]:
    """Rebuild borderless tables under a 'Table N' caption into grids. Returns
    [{'grid': [[str]], 'block_ids': set, 'y0': float}]."""
    grids: list[dict] = []
    caps = [b for b in page.blocks
            if b.type == "caption" and _TABLE_CAP_RE_DOCX.match((b.text() or "").strip())]
    for cap in caps:
        cx0, cy0, cx1, _ = cap.bbox
        col_lo, col_hi = cx0 - 25, cx0 + max(cx1 - cx0, 240) + 70
        # cells = individual lines of the table's content blocks below the caption
        cells: list[tuple] = []          # (x0, y_center, text, block_id)
        used: set[str] = set()
        for b in page.blocks:
            if b is cap or b.type in ("caption", "heading", "title", "footer", "header", "formula"):
                continue
            bxc = (b.bbox[0] + b.bbox[2]) / 2
            if not (cy0 - 4 <= b.bbox[1] <= cy0 + 230 and col_lo <= bxc <= col_hi):
                continue
            if b.translate and _is_body_prose(b):       # skip a real paragraph nearby
                continue
            for ln in b.lines:
                txt = "".join(sp.text for sp in ln.spans).strip()
                if txt and not _PAGE_NUM_RE.match(txt):
                    cells.append((ln.bbox[0], (ln.bbox[1] + ln.bbox[3]) / 2, txt, b.id))
        if len(cells) < 6:
            continue
        rows = _cluster_centers([c[1] for c in cells], gap=6.0)
        cols = _cluster_centers([c[0] for c in cells], gap=14.0)
        if len(rows) < 2 or len(cols) < 2:
            continue
        grid = [["" for _ in cols] for _ in rows]
        for x0, yc, txt, bid in cells:
            r, c = _nearest(rows, yc), _nearest(cols, x0)
            grid[r][c] = f"{grid[r][c]} {txt}".strip() if grid[r][c] else txt
            used.add(bid)
        grid = _drop_empty(grid)
        # require it to be genuinely 2-D
        if grid and len(grid) >= 2 and len(grid[0]) >= 2 \
                and sum(1 for row in grid if sum(bool(x) for x in row) >= 2) >= 2:
            grids.append({"grid": grid, "block_ids": used, "y0": cy0,
                          "x0": cx0, "x1": cx1})
    return grids


def _drop_empty(grid: list[list[str]]) -> list[list[str]]:
    """Drop all-empty rows and columns from a reconstructed grid."""
    grid = [row for row in grid if any(c.strip() for c in row)]
    if not grid:
        return grid
    keep_cols = [c for c in range(len(grid[0])) if any(row[c].strip() for row in grid)]
    return [[row[c] for c in keep_cols] for row in grid]


def _is_body_prose(blk) -> bool:
    """A running-prose paragraph (many lowercase words) — not a table row."""
    toks = (blk.text() or "").split()
    lower = sum(1 for w in toks if len(w) >= 3 and w.isalpha() and w[0].islower())
    return lower >= 5


def _col(cols: list[float], x0: float) -> int:
    # column whose left edge the block starts in (left-edge match tolerates the
    # small inset of cell text, and keeps a full-width header block in column 0)
    c = 0
    for i in range(len(cols) - 1):
        if x0 + 12 >= cols[i]:
            c = i
    return c


# --------------------------------------------------------------------------- #
# Multi-column sections + figures
# --------------------------------------------------------------------------- #
def _layout_bands(items: list, cx0: float, cx1: float) -> list:
    """Segment a page's items into vertical bands by column count, so an intra-page
    layout change is reproduced: a full-width item (title, author block, wide figure
    or table) is its own 1-column band; runs of narrow items that populate both the
    left and right half are 2-column bands. Returns [(ncol, [items])] top-to-bottom.

    Each item is ``(bbox, payload, is_wide)``; ``is_wide`` is precomputed per item
    kind (text must be near-full-width to count, figures/tables span more readily),
    so an undetected wide table cell can't fragment the body into many sections."""
    mid = (cx0 + cx1) / 2.0
    bands: list = []
    cur: list = []

    def flush() -> None:
        nonlocal cur
        if not cur:
            return
        has_left = any((it[0][0] + it[0][2]) / 2 < mid for it in cur)
        has_right = any((it[0][0] + it[0][2]) / 2 >= mid for it in cur)
        bands.append([2 if (has_left and has_right) else 1, cur])
        cur = []

    for it in sorted(items, key=lambda it: it[0][1]):     # by top y
        if it[2]:                                          # spans both columns
            flush()
            bands.append([1, [it]])
        else:
            cur.append(it)
    flush()

    # merge adjacent bands of the same column count (avoid needless section breaks)
    merged: list = []
    for nc, its in bands:
        if merged and merged[-1][0] == nc:
            merged[-1][1].extend(its)
        else:
            merged.append([nc, list(its)])
    return [(nc, its) for nc, its in merged]


def _document_two_column(doc) -> bool:
    """Whether the document is two-column overall, from the width of its body
    paragraphs (~half the content width ⇒ 2-column; near-full-width ⇒ 1-column).
    A single-column paper must NOT get 2-column bands, or a borderless table or a
    figure's stray axis labels would be split into false columns and sections."""
    ratios: list[float] = []
    for p in doc.pages:
        blks = [b for b in p.blocks if b.text().strip()]
        if not blks:
            continue
        cx0 = min(b.bbox[0] for b in blks)
        cx1 = max(b.bbox[2] for b in blks)
        cw = max(cx1 - cx0, 1.0)
        for b in blks:
            if b.type in ("paragraph", "list") and len(b.text()) > 80:
                ratios.append((b.bbox[2] - b.bbox[0]) / cw)
    if len(ratios) < 3:
        return False
    ratios.sort()
    return ratios[len(ratios) // 2] < 0.62      # median body-paragraph width


def _set_section_columns(section, n: int, page_width: float) -> None:
    """Set a docx section to n text columns (converter maps this to Hangul columns)."""
    sectPr = section._sectPr
    cols = sectPr.find(qn("w:cols"))
    if cols is None:
        cols = sectPr.makeelement(qn("w:cols"), {})
        sectPr.append(cols)
    cols.set(qn("w:num"), str(n))
    cols.set(qn("w:space"), "425")   # ~0.3" gutter


def _merge_boxes(boxes: list[tuple], gap: float = 18.0) -> list[list[float]]:
    merged = [list(b) for b in boxes]
    changed = True
    while changed:
        changed = False
        out: list[list[float]] = []
        for b in merged:
            hit = None
            for o in out:
                if not (b[0] > o[2] + gap or b[2] < o[0] - gap
                        or b[1] > o[3] + gap or b[3] < o[1] - gap):
                    hit = o
                    break
            if hit:
                hit[0] = min(hit[0], b[0]); hit[1] = min(hit[1], b[1])
                hit[2] = max(hit[2], b[2]); hit[3] = max(hit[3], b[3])
                changed = True
            else:
                out.append(b)
        merged = out
    return merged


def _overlap_frac(box, others: list) -> float:
    """Fraction of `box` area covered by any of `others` (rough, sum of overlaps)."""
    bx = (box[2] - box[0]) * (box[3] - box[1])
    if bx <= 0:
        return 0.0
    cov = 0.0
    for o in others:
        ix = max(0.0, min(box[2], o[2]) - max(box[0], o[0]))
        iy = max(0.0, min(box[3], o[3]) - max(box[1], o[1]))
        cov += ix * iy
    return min(cov / bx, 1.0)


def _region_has_ink(page, region: tuple, thresh: float = 0.010) -> bool:
    """True if a page region holds real content (a figure), not just whitespace —
    used to keep figures that are baked into a scanned page image."""
    try:
        pix = page.get_pixmap(clip=fitz.Rect(region), dpi=48, colorspace=fitz.csGRAY)
    except Exception:
        return False
    data = pix.samples
    if not data:
        return False
    dark = sum(1 for v in data if v < 190)
    return dark / len(data) > thresh


def _scan_gap_figures(page, text_boxes: list, table_boxes: list) -> list[tuple]:
    """On a scanned page the figures are part of the page image, so find the
    content-bearing vertical gaps between text/table blocks and crop those."""
    W, H = page.rect.width, page.rect.height
    if not text_boxes:
        return []
    x0 = max(0.0, min(b[0] for b in text_boxes) - 4)
    x1 = min(W, max(b[2] for b in text_boxes) + 4)
    ivs = sorted([(b[1], b[3]) for b in list(text_boxes) + list(table_boxes)])
    merged: list[list] = []
    for a, b in ivs:
        if merged and a <= merged[-1][1] + 6:
            merged[-1][1] = max(merged[-1][1], b)
        else:
            merged.append([a, b])
    figs: list[tuple] = []
    for i in range(len(merged) - 1):
        gy0, gy1 = merged[i][1], merged[i + 1][0]
        if gy1 - gy0 < 55:          # too short to hold a figure
            continue
        region = (x0, gy0 + 3, x1, gy1 - 3)
        if _region_has_ink(page, region):
            figs.append(region)
    return figs


def _figure_regions(page, text_boxes: list, table_boxes: list) -> list[tuple]:
    """Detect figure regions (vector drawings + raster images) to rasterize.

    Papers draw figures as vector graphics, so we cluster drawing/image rects and
    keep sizeable clusters that are not mostly text and not a table."""
    W, H = page.rect.width, page.rect.height
    area = W * H
    # Scanned page: a near-full-page image with OCR text over it. Its figures are baked
    # into the scan, so crop the content-bearing gaps between text rather than embed
    # the whole scan (which would drop the translation and show the original English).
    for im in page.get_images(full=True):
        with contextlib.suppress(Exception):
            r = page.get_image_bbox(im)
            if r.width * r.height > 0.85 * area and len(text_boxes) >= 3:
                return _scan_gap_figures(page, text_boxes, table_boxes)
    rects: list[tuple] = []
    for d in page.get_drawings():
        r = d.get("rect")
        if r and r.width > 3 and r.height > 3 and r.width < W and r.height < H:
            rects.append((r.x0, r.y0, r.x1, r.y1))
    for im in page.get_images(full=True):
        with contextlib.suppress(Exception):
            r = page.get_image_bbox(im)
            if r.width > 3 and r.height > 3:
                rects.append((r.x0, r.y0, r.x1, r.y1))
    if not rects:
        return []
    figs = []
    for c in _merge_boxes(rects, gap=18.0):
        w, h = c[2] - c[0], c[3] - c[1]
        if w < 45 or h < 45 or w * h < 0.02 * area:
            continue
        if _overlap_frac(c, text_boxes) > 0.55:     # a text area, not a figure
            continue
        if _overlap_frac(c, table_boxes) > 0.5:      # a table, handled elsewhere
            continue
        figs.append(tuple(c))
    return figs


def _expand_figs_with_labels(figs: list, blocks, page_w: float) -> list:
    """Grow each figure box to swallow the small labels around it (axis ticks,
    numbers, legends) so they are rasterized into the image instead of leaking out
    as stray text. Only small, short, non-caption blocks close to the box qualify —
    never a caption or a body paragraph."""
    grown = [list(f) for f in figs]
    for g in grown:
        changed = True
        while changed:
            changed = False
            for b in blocks:
                if b.type == "caption" or not b.text().strip():
                    continue
                bx0, by0, bx1, by1 = b.bbox
                if (bx1 - bx0) > 0.22 * page_w or len(b.text()) > 30:
                    continue                       # a real paragraph, not a label
                m = 26.0
                if bx0 >= g[0] - m and bx1 <= g[2] + m and by0 >= g[1] - m and by1 <= g[3] + m:
                    if not (g[0] <= bx0 and g[2] >= bx1 and g[1] <= by0 and g[3] >= by1):
                        g[0] = min(g[0], bx0); g[1] = min(g[1], by0)
                        g[2] = max(g[2], bx1); g[3] = max(g[3], by1)
                        changed = True
    return [tuple(g) for g in grown]


def _tu_by_block(doc: Document, tr: TranslationResult) -> dict:
    tu_by_id = tr.tu_map()
    out: dict = {}
    for tu in (doc.translation_units or tr.translation_units):
        cur = tu_by_id.get(tu.id, tu)
        for bid in tu.block_ids:
            out[bid] = cur
    return out


def _clean(text: str) -> str:
    return _strip_math_delims(text or "").strip()


def _emit(d, style_type: str, text: str) -> None:
    text = _clean(text)
    if not text:
        return
    if style_type == "title":
        d.add_heading(text, level=0)
    elif style_type == "heading":
        level = 2 if _SUBHEAD_RE.match(text) else 1
        d.add_heading(text, level=level)
    else:  # paragraph / list / formula — flowing body text (marker kept in text)
        d.add_paragraph(text)


def _emit_table(d, spec: dict, grid: list[list[str]]) -> None:
    nrows, ncols = len(grid), len(grid[0])
    table = d.add_table(rows=nrows, cols=ncols)
    with contextlib.suppress(KeyError):
        table.style = "Table Grid"          # visible borders
    for r in range(nrows):
        for c in range(ncols):
            if grid[r][c]:
                table.cell(r, c).text = grid[r][c]


def _emit_figure(d, pdf_page, bbox: tuple, tmpdir: Path, idx: int, avail_pt: float) -> None:
    """Rasterize a figure region (vector or raster) and embed it, centered."""
    try:
        pix = pdf_page.get_pixmap(clip=fitz.Rect(bbox), dpi=150)
        png = tmpdir / f"fig_{pdf_page.number}_{idx}.png"
        pix.save(str(png))
    except Exception:
        return
    w = min(bbox[2] - bbox[0], avail_pt)
    p = d.add_paragraph()
    p.alignment = 1  # center
    with contextlib.suppress(Exception):
        p.add_run().add_picture(str(png), width=Pt(w))


def export_docx(doc: Document, tr: TranslationResult, out_path: Path,
                pdf_path: str | None = None) -> Path:
    """Build a flowing .docx from the translated document.

    Reconstructs ruled tables, applies the page's column count (1- or 2-column
    sections), and embeds figures (rasterized from the PDF, vector or raster).
    ``pdf_path`` (the job's current PDF) drives tables/figures/columns; it falls
    back to the IR's ``source_path``. Without a PDF, tables/figures degrade out.
    """
    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    d = Docx()
    try:
        normal = d.styles["Normal"]
        normal.font.name = "맑은 고딕"
        normal.font.size = Pt(11)
    except Exception:
        pass

    tu_by_block = _tu_by_block(doc, tr)

    pdf = None
    for cand in (pdf_path, doc.source_path):
        if cand and Path(cand).exists():
            try:
                pdf = fitz.open(cand)
                break
            except Exception:
                pdf = None

    def blk_text(blk) -> str:
        if not blk.translate:
            return _clean(blk.text())
        tu = tu_by_block.get(blk.id)
        if tu is None:
            return _clean(blk.text())
        done = tu.status == "done" and (tu.target_text or "").strip()
        return _clean(tu.target_text if done else blk.text())

    two_col = _document_two_column(doc)
    tmpdir = Path(tempfile.mkdtemp(prefix="paperko_fig_"))
    cur_ncol: int | None = None
    emitted_tu: set[str] = set()
    first_page = True
    try:
        for page in sorted(doc.pages, key=lambda p: p.index):
            pdf_page = pdf.load_page(page.index) if (pdf and page.index < pdf.page_count) else None

            # --- tables: detect + assign cells (same as before) ---
            specs: list[dict] = []
            if pdf_page is not None:
                with contextlib.suppress(Exception):
                    specs = _page_tables(pdf_page)
            assign: dict[str, tuple[int, int, int]] = {}
            grids: list[list[list[str]]] = [
                [["" for _ in range(len(s["cols"]) - 1)] for _ in range(len(s["rows"]) - 1)]
                for s in specs
            ]
            for si, s in enumerate(specs):
                for blk in page.blocks:
                    if blk.type == "formula" or _PAGE_NUM_RE.match((blk.text() or "").strip()):
                        continue
                    xc = (blk.bbox[0] + blk.bbox[2]) / 2
                    yc = (blk.bbox[1] + blk.bbox[3]) / 2
                    if not (s["x0"] - 2 <= xc <= s["x1"] + 2 and s["y0"] - 2 <= yc <= s["y1"] + 2):
                        continue
                    r = _band(s["rows"], yc)
                    if r is None:
                        continue
                    c = _col(s["cols"], blk.bbox[0])
                    assign[blk.id] = (si, r, c)
                    txt = blk_text(blk)
                    if txt:
                        cell = grids[si][r][c]
                        grids[si][r][c] = f"{cell} {txt}".strip() if cell else txt
            dead: set[int] = set()
            for si, g in enumerate(grids):
                rows_hit = {r for r, row in enumerate(g) if any(row)}
                cols_hit = {c for row in g for c, cell in enumerate(row) if cell}
                if len(rows_hit) < 2 or len(cols_hit) < 2:
                    dead.add(si)
            table_boxes = [(s["x0"], s["y0"], s["x1"], s["y1"])
                           for si, s in enumerate(specs) if si not in dead]

            # --- figures ---
            figs: list[tuple] = []
            if pdf_page is not None:
                with contextlib.suppress(Exception):
                    figs = _figure_regions(pdf_page, [b.bbox for b in page.blocks], table_boxes)
                    if figs:
                        figs = _expand_figs_with_labels(figs, page.blocks, page.width)

            def in_fig(blk) -> bool:
                xc = (blk.bbox[0] + blk.bbox[2]) / 2
                yc = (blk.bbox[1] + blk.bbox[3]) / 2
                return any(f[0] - 2 <= xc <= f[2] + 2 and f[1] - 2 <= yc <= f[3] + 2 for f in figs)

            # --- borderless tables rebuilt from cell positions (no ruling lines) ---
            bgrids = _borderless_grids(page)
            bconsumed: set[str] = set()
            for g in bgrids:
                bconsumed |= g["block_ids"]

            # --- collect page items (bbox, payload, is_wide) ---
            raw: list[tuple] = []   # (bbox, payload, kind_is_figtbl)
            for si, s in enumerate(specs):
                if si not in dead:
                    raw.append(((s["x0"], s["y0"], s["x1"], s["y1"]), ("tbl", si), True))
            for fi, f in enumerate(figs):
                raw.append((tuple(f), ("fig", fi, f), True))
            for gi, g in enumerate(bgrids):
                # place the reconstructed table just below its caption
                y = g["y0"] + 1
                raw.append(((g["x0"], y, g["x1"], y + 10), ("btbl", gi, g), True))
            for blk in page.blocks:
                if blk.id in assign and assign[blk.id][0] not in dead:
                    continue                                   # inside a real table
                if blk.id in bconsumed:
                    continue                                   # inside a borderless table
                if in_fig(blk):
                    continue                                   # baked into a figure image
                if blk.type == "formula" and _PAGE_NUM_RE.match((blk.text() or "").strip()):
                    continue
                raw.append((tuple(blk.bbox), ("blk", blk), False))
            if not raw:
                continue

            # A band divider must genuinely span the page. Figures/tables span at
            # 0.62 of the content width; text must be near-full-width (0.85), so a
            # wide table-cell remnant or a long line can't fragment the 2-column body.
            cx0 = min(it[0][0] for it in raw)
            cx1 = max(it[0][2] for it in raw)
            cw = max(cx1 - cx0, 1.0)
            mid = (cx0 + cx1) / 2.0
            items = [
                (bb, payload, (bb[2] - bb[0]) >= (0.62 if figtbl else 0.85) * cw)
                for bb, payload, figtbl in raw
            ]
            # Single-column document: everything is one 1-column band (in reading
            # order), so tables / figure labels are never split into false columns.
            if not two_col:
                bands = [(1, items)]
            else:
                # First page: everything above the "1. Introduction" heading (title,
                # authors, highlights, abstract) is forced into one 1-column band so
                # it is not split into the body's two columns.
                bands = None
                if first_page:
                    cutoff = _frontmatter_cutoff(page)
                    if cutoff is not None:
                        # tolerance so the body's first right-column line (level with
                        # the heading) stays in the 2-column body, not the front matter.
                        front = [it for it in items if it[0][1] < cutoff - 8]
                        rest = [it for it in items if it[0][1] >= cutoff - 8]
                        if front:
                            bands = [(1, front), *_layout_bands(rest, cx0, cx1)]
                if bands is None:
                    bands = _layout_bands(items, cx0, cx1)
            first_page = False

            for bncol, bitems in bands:
                if cur_ncol is None:
                    _set_section_columns(d.sections[0], bncol, page.width)
                    cur_ncol = bncol
                elif bncol != cur_ncol:
                    _set_section_columns(d.add_section(WD_SECTION.CONTINUOUS), bncol, page.width)
                    cur_ncol = bncol
                avail = (page.width - 144.0) if bncol == 1 else (page.width - 144.0 - 24.0) / 2.0

                def _key(it):
                    bb = it[0]
                    col = 0 if bncol == 1 else (0 if (bb[0] + bb[2]) / 2 < mid else 1)
                    return (col, bb[1])

                for it in sorted(bitems, key=_key):
                    payload = it[1]
                    if payload[0] == "tbl":
                        _emit_table(d, specs[payload[1]], grids[payload[1]])
                    elif payload[0] == "btbl":
                        _emit_table(d, None, payload[2]["grid"])
                    elif payload[0] == "fig":
                        _emit_figure(d, pdf_page, payload[2], tmpdir, payload[1], avail)
                    else:
                        blk = payload[1]
                        tu = tu_by_block.get(blk.id)
                        if blk.translate and tu is not None:
                            if tu.id in emitted_tu:
                                continue
                            emitted_tu.add(tu.id)
                        _emit(d, blk.type, blk_text(blk))

        if pdf is not None:
            pdf.close()
        d.save(str(out_path))
    finally:
        with contextlib.suppress(Exception):
            import shutil
            shutil.rmtree(tmpdir, ignore_errors=True)
    return out_path


_convert_lock = threading.Lock()
_convert_mod = None  # cached bundled converter module


def _load_converter():
    """Import the bundled docx→hwpx converter in-process (cached).

    Done in-process rather than as a subprocess so it works no matter how the
    engine is packaged (embeddable Python or a frozen exe, where sys.executable
    is not a plain interpreter) and never flashes a console window on Windows.
    """
    global _convert_mod
    if _convert_mod is not None:
        return _convert_mod
    scripts = Path(__file__).parent / "hwpx" / "scripts"
    if not (scripts / "convert.py").exists():
        raise FileNotFoundError(f"bundled hwpx converter missing: {scripts / 'convert.py'}")
    sp = str(scripts)
    if sp not in sys.path:
        sys.path.insert(0, sp)
    importlib.import_module("omml_to_hwp")   # sibling import used by convert.py
    _convert_mod = importlib.import_module("convert")
    return _convert_mod


def _docx_to_hwpx(docx_path: Path, out_path: Path) -> Path:
    """Convert a .docx to .hwpx using the bundled Hancom converter (offline)."""
    skeleton = Path(__file__).parent / "hwpx" / "assets" / "default_skeleton.hwpx"
    # Serialize (the converter module keeps some module-level state) and send any
    # of its stdout to stderr so it never corrupts the JSON-RPC channel on stdout.
    with _convert_lock:
        conv = _load_converter()
        with contextlib.redirect_stdout(sys.stderr):
            conv.convert(str(docx_path), str(skeleton), str(out_path), verbose=False)
    if not out_path.exists():
        raise RuntimeError("hwpx conversion produced no output")
    return out_path


def export_document(doc: Document, tr: TranslationResult, out_path: Path, fmt: str,
                    pdf_path: str | None = None) -> Path:
    """Export the translated document as ``docx`` or ``hwpx``."""
    out_path = Path(out_path)
    fmt = (fmt or "").lower()
    if fmt == "docx":
        return export_docx(doc, tr, out_path, pdf_path)
    if fmt == "hwpx":
        tmp_docx = out_path.with_name(out_path.stem + ".__tmp__.docx")
        try:
            export_docx(doc, tr, tmp_docx, pdf_path)
            _docx_to_hwpx(tmp_docx, out_path)
        finally:
            try:
                tmp_docx.unlink()
            except OSError:
                pass
        return out_path
    raise ValueError(f"unsupported export format: {fmt!r} (want 'docx' or 'hwpx')")
