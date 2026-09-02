import zipfile

import pytest
from docx import Document as Docx

from translate_engine.docx_export import export_document
from translate_engine.ir import (
    Block, Document, Line, Page, Span, TranslationResult, TU,
)


def _blk(bid, typ, text, order):
    b = Block(id=bid, page=0, type=typ, bbox=[0, 0, 100, 20], reading_order=order)
    b.lines = [Line(bbox=[0, 0, 100, 10], spans=[Span(text=text, bbox=[0, 0, 100, 10])])]
    b.translate = typ != "formula"
    return b


def _doc_and_tr():
    blocks = [
        _blk("b1", "heading", "제목", 0),
        _blk("b2", "paragraph", "원문 문단", 1),
        _blk("b3", "formula", "- 2 -", 2),   # page number → must be dropped
    ]
    doc = Document(doc_id="d", source_path="", meta={},
                   pages=[Page(index=0, width=595, height=842, blocks=blocks)])
    tr = TranslationResult(doc_id="d", translation_units=[
        TU(id="t1", block_ids=["b1"], source_text="제목", status="done", target_text="Title"),
        TU(id="t2", block_ids=["b2"], source_text="원문 문단", status="done",
           target_text="A translated paragraph."),
    ])
    doc.translation_units = tr.translation_units
    return doc, tr


def test_export_docx_uses_translation_and_drops_page_numbers(tmp_path):
    doc, tr = _doc_and_tr()
    out = export_document(doc, tr, tmp_path / "o.docx", "docx")
    d = Docx(str(out))
    texts = [p.text for p in d.paragraphs if p.text.strip()]
    assert "Title" in texts                     # heading translated
    assert "A translated paragraph." in texts   # body translated
    assert "- 2 -" not in " ".join(texts)       # page number dropped
    styles = {p.text: p.style.name for p in d.paragraphs if p.text.strip()}
    assert styles["Title"].startswith("Heading")


def test_export_hwpx_produces_valid_hancom_zip(tmp_path):
    doc, tr = _doc_and_tr()
    out = export_document(doc, tr, tmp_path / "o.hwpx", "hwpx")
    assert out.exists()
    with zipfile.ZipFile(out) as z:
        names = z.namelist()
        assert "mimetype" in names
        assert any(n.startswith("Contents/section") for n in names)
        body = z.read("Contents/section0.xml").decode("utf-8", "ignore")
        assert "A translated paragraph." in body


def test_borderless_grid_reconstruction():
    from translate_engine.ir import Line, Span
    from translate_engine.docx_export import _borderless_grids

    def cell(text, x, y):
        return Line(bbox=[x, y, x + 30, y + 10],
                    spans=[Span(text=text, bbox=[x, y, x + 30, y + 10])])

    cap = Block(id="c", page=0, type="caption", bbox=[40, 100, 200, 112])
    cap.lines = [cell("Table 3 Load on structure", 40, 100)]
    # a table: header row + 2 data rows, cells at 3 column x-positions
    rowblk = Block(id="t", page=0, type="paragraph", bbox=[40, 120, 300, 160])
    rowblk.translate = False
    rowblk.lines = [
        cell("Member", 40, 122), cell("G1", 120, 122), cell("G2", 200, 122),
        cell("Dead", 40, 138), cell("4.75", 120, 138), cell("6.71", 200, 138),
        cell("Live", 40, 154), cell("3.42", 120, 154), cell("4.82", 200, 154),
    ]
    page = Page(index=0, width=595, height=842, blocks=[cap, rowblk])
    grids = _borderless_grids(page)
    assert len(grids) == 1
    grid = grids[0]["grid"]
    assert grid[0] == ["Member", "G1", "G2"]
    assert grid[1] == ["Dead", "4.75", "6.71"]
    assert grid[2] == ["Live", "3.42", "4.82"]


def test_export_rejects_unknown_format(tmp_path):
    doc, tr = _doc_and_tr()
    with pytest.raises(ValueError):
        export_document(doc, tr, tmp_path / "o.txt", "txt")
